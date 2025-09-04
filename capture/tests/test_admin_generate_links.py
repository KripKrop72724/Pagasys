import io
from datetime import timedelta

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.urls import reverse
from django.utils import timezone
from django.test import TestCase
from openpyxl import load_workbook

from capture.models import EnrollmentLink, FaceEnrollment
from pagasys.tests.test_models import ModelFactoryMixin


class GenerateEnrollmentLinksAdminTests(ModelFactoryMixin, TestCase):
    def setUp(self):
        call_command("initgroups", verbosity=0)
        self.company = self.create_company()
        self.branch1 = self.create_branch(self.company, "B1")
        self.branch2 = self.create_branch(self.company, "B2")
        self.department1 = self.create_department(self.branch1, "D1")
        self.department2 = self.create_department(self.branch2, "D2")
        self.license = self.create_license(
            company=self.company,
            branches=[self.branch1, self.branch2],
            license_no="LIC1",
            max_visas=10,
        )
        User = get_user_model()
        self.admin = User.objects.create_superuser(
            username="admin",
            password="pass",
            trade_license=self.license,
            department=self.department1,
            hire_date="2024-01-01",
            employment_type="permanent",
            visa_type="company",
        )
        self.client.force_login(self.admin)
        self.client.defaults["HTTP_HOST"] = "0.0.0.0"
        self.emp1 = User.objects.create_user(
            username="e1",
            password="pass",
            first_name="E1",
            last_name="User",
            trade_license=self.license,
            department=self.department1,
            hire_date="2024-01-01",
            employment_type="permanent",
            visa_type="company",
        )
        self.emp2 = User.objects.create_user(
            username="e2",
            password="pass",
            first_name="E2",
            last_name="User",
            trade_license=self.license,
            department=self.department2,
            hire_date="2024-01-01",
            employment_type="permanent",
            visa_type="company",
        )

    def test_generate_links_downloads_excel(self):
        url = reverse("admin:capture_faceenrollment_generate_links")
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertEqual(EnrollmentLink.objects.count(), 3)
        link = EnrollmentLink.objects.get(employee=self.emp1)
        self.assertEqual(link.max_uses, 1)
        self.assertAlmostEqual(
            link.expires_at, timezone.now() + timedelta(days=2), delta=timedelta(seconds=5)
        )
        wb = load_workbook(io.BytesIO(resp.content))
        self.assertIn("B1", wb.sheetnames)
        self.assertIn("B2", wb.sheetnames)
        ws = wb["B1"]
        self.assertEqual(ws["A1"].value, "Name")
        self.assertEqual(ws["B1"].value, "Enrollment URL")
        self.assertEqual(ws.max_row, 3)
        self.assertTrue(ws["B2"].value.startswith("http"))

    def test_generate_links_skips_enrolled_employees(self):
        FaceEnrollment.objects.create(
            employee=self.emp1,
            collection_id="c",
            face_ids=["f1"],
            status="active",
        )
        url = reverse("admin:capture_faceenrollment_generate_links")
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(EnrollmentLink.objects.count(), 2)
        employee_ids = set(EnrollmentLink.objects.values_list("employee", flat=True))
        self.assertIn(self.admin.id, employee_ids)
        self.assertIn(self.emp2.id, employee_ids)
        self.assertNotIn(self.emp1.id, employee_ids)

        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb["B1"]
        self.assertEqual(ws.max_row, 2)  # header + one row
        names = [ws[f"A{row}"].value for row in range(2, ws.max_row + 1)]
        self.assertNotIn(self.emp1.get_full_name(), names)
