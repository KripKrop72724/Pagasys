import itertools
import tempfile

import pytest
from django.contrib.messages import get_messages
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from pagasys.models import Employee
from capture.models import FaceEnrollment

from .test_views import (
    company,
    branch,
    branch2,
    department,
    department2,
    client,
)  # reuse fixtures


@pytest.mark.django_db
def test_command_rebuilds_and_skips(monkeypatch, department, department2, capsys):
    e1 = Employee.objects.create_user(
        username="emp1",
        password="pw",
        hire_date=timezone.localdate(),
        employment_type="permanent",
        department=department,
        visa_type="personal",
        first_name="A",
    )
    e2 = Employee.objects.create_user(
        username="emp2",
        password="pw",
        hire_date=timezone.localdate(),
        employment_type="permanent",
        department=department2,
        visa_type="personal",
        first_name="B",
    )

    from capture.management.commands import rebuild_face_enrollments as cmd

    def fake_fetch(company_id, employee_id):
        if employee_id == e1.id:
            return [b"1", b"2", b"3"]
        return [b"1", b"2", b"3", b"4"]

    counter = itertools.count()

    def fake_index(company_id, employee_id, img):
        return [f"id{next(counter)}"]

    monkeypatch.setattr(cmd, "fetch_enroll_images", fake_fetch)
    monkeypatch.setattr(cmd, "index_faces", fake_index)

    command = cmd.Command()
    path = command.handle()
    skipped = command.skipped

    assert skipped == [e1.get_full_name()]
    assert not FaceEnrollment.objects.filter(employee=e1).exists()
    fe2 = FaceEnrollment.objects.get(employee=e2)
    assert len(fe2.face_ids) == 4

    wb = load_workbook(path)
    ws_re = wb["Reenrolled"]
    assert ws_re.max_row == 2
    assert ws_re["A2"].value == e2.branch.name
    assert ws_re["B2"].value == e2.get_full_name()
    ws_sk = wb["Skipped"]
    assert ws_sk.max_row == 2
    assert ws_sk["A2"].value == e1.get_full_name()

    out = capsys.readouterr().out
    assert e1.get_full_name() in out


@pytest.mark.django_db
def test_admin_view_streams_workbook(monkeypatch, client, department):
    admin = Employee.objects.create_superuser(
        username="admin",
        password="pw",
        hire_date=timezone.localdate(),
        employment_type="permanent",
        department=department,
        visa_type="personal",
    )
    client.force_login(admin)

    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()

    class FakeCommand:
        def __init__(self):
            self.skipped = ["Foo"]

        def handle(self):
            return tmp.name

    monkeypatch.setattr("capture.admin.Command", FakeCommand)

    resp = client.get(reverse("admin:capture_faceenrollment_rebuild_enrollments"))
    assert resp.status_code == 200
    assert resp["Content-Type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    messages = list(get_messages(resp.wsgi_request))
    assert "Foo" in messages[0].message
