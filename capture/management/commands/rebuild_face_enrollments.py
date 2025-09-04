from __future__ import annotations

import tempfile
from typing import Any, Tuple

from django.core.management.base import BaseCommand
from openpyxl import Workbook

from pagasys.models import Employee
from capture.aws import (
    fetch_enroll_images,
    index_faces,
    company_collection_id,
)
from capture.models import FaceEnrollment


def rebuild_from_s3() -> Tuple[str, list[str]]:
    reenrolled: list[tuple[str, str]] = []
    skipped: list[str] = []

    for emp in Employee.objects.all():
        if emp.company is None:
            continue
        images = fetch_enroll_images(emp.company.id, emp.id)
        if len(images) < 4:
            skipped.append(emp.get_full_name())
            continue
        face_ids: list[str] = []
        for img in images:
            face_ids.extend(index_faces(emp.company.id, emp.id, img))
        FaceEnrollment.objects.update_or_create(
            employee=emp,
            defaults={
                "collection_id": company_collection_id(emp.company.id),
                "face_ids": face_ids,
                "status": "active",
            },
        )
        branch_name = emp.branch.name if emp.branch else "Unknown"
        reenrolled.append((branch_name, emp.get_full_name()))

    wb = Workbook()
    ws_re = wb.create_sheet("Reenrolled")
    ws_sk = wb.create_sheet("Skipped")
    wb.remove(wb.active)

    ws_re.append(["Branch", "Employee"])
    for branch, name in reenrolled:
        ws_re.append([branch, name])

    ws_sk.append(["Employee"])
    for name in skipped:
        ws_sk.append([name])

    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    wb.save(tmp.name)
    tmp.close()
    return tmp.name, skipped


class Command(BaseCommand):
    help = "Rebuild face enrollments from images stored in S3"

    def handle(self, *args: Any, **options: Any):
        path, skipped = rebuild_from_s3()
        self.skipped = skipped
        if skipped:
            self.stdout.write("Skipped employees: " + ", ".join(skipped))
        else:
            self.stdout.write("Skipped employees: none")
        return path
