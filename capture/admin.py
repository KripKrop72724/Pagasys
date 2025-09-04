import io
from collections import defaultdict
from datetime import timedelta

from django.contrib import admin
from django.http import HttpResponse
from django.shortcuts import redirect
from django.urls import path
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font

from pagasys.models import Employee
from pagasys.utils import scope_queryset
from .models import AttendanceDevice, FaceEnrollment, EnrollmentLink, PunchEvent, PunchException
from .aws import count_all_faces, delete_all_faces


class ScopedAdminMixin:
    def get_queryset(self, request):
        return scope_queryset(super().get_queryset(request), request.user)


@admin.register(AttendanceDevice)
class AttendanceDeviceAdmin(ScopedAdminMixin, admin.ModelAdmin):
    """Admin interface for configuring capture devices and geofence settings."""
    list_display = [
        "name",
        "company",
        "branch",
        "department",
        "project",
        "is_active",
        "last_seen",
        "latitude",
        "longitude",
        "radius_m",
        "note",
    ]
    list_filter = ["company", "branch", "department", "project", "is_active"]
    search_fields = ["name", "api_key"]


@admin.register(FaceEnrollment)
class FaceEnrollmentAdmin(ScopedAdminMixin, admin.ModelAdmin):
    """View and manage employee face enrollments."""
    list_display = ["employee", "status", "created_at", "updated_at"]
    actions = ["clear_aws_faces"]
    change_list_template = "admin/capture/faceenrollment/change_list.html"

    def get_urls(self):
        urls = super().get_urls()
        my_urls = [
            path(
                "clear-all-aws-faces/",
                self.admin_site.admin_view(self.clear_all_aws_faces),
                name="capture_faceenrollment_clear_all_aws_faces",
            ),
            path(
                "generate-links/",
                self.admin_site.admin_view(self.generate_links),
                name="capture_faceenrollment_generate_links",
            ),
        ]
        return my_urls + urls

    def changelist_view(self, request, extra_context=None):
        extra_context = extra_context or {}
        try:
            extra_context["aws_face_total"] = count_all_faces()
        except RuntimeError:
            extra_context["aws_face_total"] = 0
        return super().changelist_view(request, extra_context=extra_context)

    @admin.action(description="Clear selected face enrollments")
    def clear_aws_faces(self, request, queryset):
        count = 0
        for fe in queryset:
            fe.delete()
            count += 1
        self.message_user(request, f"Cleared {count} enrollment(s)")

    def clear_all_aws_faces(self, request):
        delete_all_faces()
        FaceEnrollment.objects.all().delete()
        self.message_user(request, "Cleared all face enrollments")
        return redirect("..")

    def generate_links(self, request):
        employees = scope_queryset(Employee.objects.all(), request.user)
        now = timezone.now()
        expires = now + timedelta(days=2)
        links_by_branch: dict[str, list[tuple[str, str]]] = defaultdict(list)
        for emp in employees:
            link = EnrollmentLink.objects.create(
                employee=emp,
                expires_at=expires,
                max_uses=1,
            )
            branch = emp.branch.name if emp.branch else "Unknown"
            links_by_branch[branch].append((emp.get_full_name(), link.url))

        wb = Workbook()
        wb.remove(wb.active)
        for branch, rows in links_by_branch.items():
            ws = wb.create_sheet(title=branch[:31])
            ws.append(["Name", "Enrollment URL"])
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for name, url in rows:
                ws.append([name, url])
            for column_cells in ws.columns:
                length = max(len(str(cell.value)) if cell.value else 0 for cell in column_cells)
                ws.column_dimensions[column_cells[0].column_letter].width = length + 2
            ws.freeze_panes = "A2"

        with io.BytesIO() as bio:
            wb.save(bio)
            bio.seek(0)
            response = HttpResponse(
                bio.getvalue(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        response["Content-Disposition"] = "attachment; filename=face_enrollment_links.xlsx"
        return response



@admin.register(EnrollmentLink)
class EnrollmentLinkAdmin(ScopedAdminMixin, admin.ModelAdmin):
    """Inspect enrollment links issued to employees.

    Tokens are generated automatically and are immutable. The full enrollment
    URL is displayed so it can be copied directly from the admin.
    """

    list_display = ["employee", "token", "url", "expires_at", "used_at", "uses", "max_uses"]
    readonly_fields = ["token", "url", "uses", "used_at"]
    fields = ["employee", "token", "url", "expires_at", "max_uses", "uses", "used_at"]

    @admin.display(description="Enrollment URL")
    def url(self, obj):
        return obj.url


@admin.register(PunchEvent)
class PunchEventAdmin(ScopedAdminMixin, admin.ModelAdmin):
    """Review raw punch events and their validation state."""
    list_display = [
        "company",
        "device",
        "employee",
        "action",
        "device_ts",
        "face_matched",
        "face_confidence",
        "requires_face",
        "out_of_scope",
        "geofence_ok",
        "geofence_rule_violation",
        "roster_fallback",
        "roster_date",
        "external_id",
    ]
    list_filter = [
        "company",
        "device",
        "face_matched",
        "requires_face",
        "out_of_scope",
        "geofence_ok",
        "geofence_rule_violation",
        "roster_fallback",
        "roster_date",
    ]
    search_fields = ["employee__username", "device__name", "external_id"]


@admin.register(PunchException)
class PunchExceptionAdmin(ScopedAdminMixin, admin.ModelAdmin):
    """Display punch events that triggered exceptions."""
    list_display = ["event", "kind", "details"]
    list_filter = ["kind"]
