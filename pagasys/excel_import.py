from __future__ import annotations

import re

from datetime import datetime, date
from typing import IO, List, Dict, Tuple

from django.contrib.auth.models import Group
from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook

from .models import Branch, Department, Designation, Employee, TradeLicense
from .policy.permissions import EMPLOYEE_ROLE

# Mapping of provided nationality names to ISO country codes
NATIONALITY_MAP = {
    "BANGLADESH": "BD",
    "EGYPTION": "EG",
    "EGYPT": "EG",
    "EGYPTIAN": "EG",
    "INDIA": "IN",
    "INDIAN": "IN",
    "NEPAL": "NP",
    "NEPALI": "NP",
    "PAKISTAN": "PK",
    "PAKISTANI": "PK",
    "SRI LANKA": "LK",
    "SRI LANKAN": "LK",
    "GHANA": "GH",
    "NIGERIAN": "NG",
}


def _normalize_header(value: object) -> str:
    return "".join(ch for ch in str(value or "").upper() if ch.isalnum())


MAIN_FORMAT_HEADER_MAP = {
    "FIRSTNAME": "first_name",
    "MIDDLENAME": "middle_name",
    "LASTNAME": "last_name",
    "BRANCH": "branch",
    "TRADELICENSE": "trade_license",
    "TRADERLICENSE": "trade_license",
    "TRADRLICENSE": "trade_license",
    "TRADELICENCE": "trade_license",
    "TRADERLICENCE": "trade_license",
    "TRADRLICENCE": "trade_license",
    "TRADLICENSE": "trade_license",
    "VISATYPE": "visa_type",
    "PAYMENTSTATUS": "payment_status",
    "WPSACCOUNTNUMBER": "wps_account_number",
    "WPSACCNUMBER": "wps_account_number",
    "WPSACNUMBER": "wps_account_number",
    "WPSACCNUMEBR": "wps_account_number",
    "WPSACCOUNTNO": "wps_account_number",
    "WPSACCNO": "wps_account_number",
    "DESIGNATION": "designation",
    "C3NUMBER": "c3_number",
    "C3NO": "c3_number",
    "C3ID": "c3_number",
    "HIREDATE": "hire_date",
    "DATEOFJOINING": "hire_date",
    "DOJ": "hire_date",
    "GENDER": "gender",
}


MAIN_FORMAT_REQUIRED_COLUMNS = {
    "first_name",
    "branch",
    "trade_license",
    "visa_type",
    "payment_status",
    "hire_date",
}


VISA_TYPE_MAP = {
    "COMPANY": "company",
    "COMPANYVISA": "company",
    "PERSONAL": "personal",
    "PERSONALVISA": "personal",
    "OWN": "personal",
    "OWNVISA": "personal",
    "VISIT": "visit",
    "VISITVISA": "visit",
}


PAYMENT_STATUS_MAP = {
    "WPS": "wps",
    "CASH": "cash",
}


def _stringify(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    return str(value)


def _is_empty(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def _parse_hire_date(raw: object) -> date:
    if raw is None:
        raise ValueError("missing hire date")
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw).strip()
    if not text:
        raise ValueError("missing hire date")
    for fmt in ("%Y/%m/%d", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    try:
        return date.fromisoformat(text)
    except ValueError:
        pass
    raise ValueError(f"could not parse hire date '{text}'")


def _normalize_gender(raw: object) -> str:
    text = _stringify(raw).strip().lower()
    if not text:
        return ""
    if text in {"male", "m"}:
        return "male"
    if text in {"female", "f"}:
        return "female"
    raise ValueError(f"Unsupported gender '{raw}'")


def _split_name(raw: str) -> Tuple[str, str, str]:
    parts = raw.split()
    first = parts[0].title()
    if len(parts) == 1:
        return first, "", ""
    if len(parts) == 2:
        return first, "", parts[1].title()
    # more than 2 parts
    middle = parts[1].title()
    last = parts[2].title()
    return first, middle, last


def import_employee_workbook(file: IO[bytes]) -> Dict[str, object]:
    """Import employees from an Excel workbook.

    Returns a dict with counts of created employees and any error messages.
    """

    wb = load_workbook(file, data_only=True)
    group, _ = Group.objects.get_or_create(name=EMPLOYEE_ROLE)
    results = {"created": 0, "errors": []}  # type: Dict[str, object]

    sheets = [("OWN", "personal"), ("VISIT", "visit")]
    for sheet_name, visa_type in sheets:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        headers = [cell.value for cell in ws[1]]
        header_index = {str(h).strip(): i for i, h in enumerate(headers)}
        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            try:
                name = str(row[header_index["NAME"]]).strip()
                first, middle, last = _split_name(name)

                designation_name = str(row[header_index["DESIGNATION"]]).strip()
                branch_name = str(row[header_index["WORKING BRANCH"]]).strip()
                branch = Branch.objects.filter(name__iexact=branch_name).first()
                if not branch:
                    raise ValueError(f"Unknown branch '{branch_name}'")
                department, _ = Department.objects.get_or_create(branch=branch, name="General")

                designation = None
                if designation_name:
                    designation, _ = Designation.objects.get_or_create(
                        company=branch.company, name=designation_name
                    )

                nationality_raw = str(row[header_index["NATIONALITY"]]).strip().upper()
                country_code = NATIONALITY_MAP.get(nationality_raw)
                if not country_code:
                    raise ValueError(f"Unsupported nationality '{nationality_raw}'")

                hire_raw = row[header_index["DATE OF JOINING"]]
                hire_date: date | None = None
                if hire_raw:
                    if isinstance(hire_raw, datetime):
                        hire_date = hire_raw.date()
                    elif isinstance(hire_raw, date):
                        hire_date = hire_raw
                    else:
                        hire_date = datetime.strptime(str(hire_raw), "%Y/%m/%d").date()

                tp_number = str(row[header_index["TP NUMBER"]]).strip()
                username_base = "".join(name.split()).lower()
                username = username_base
                suffix = 1
                while Employee.objects.filter(username=username).exists():
                    username = f"{username_base}{suffix}"
                    suffix += 1

                special_notes = ""
                if sheet_name == "VISIT":
                    special_notes = str(row[header_index.get("VISA STATUS")]).strip()

                gender_raw = str(row[header_index["GENDER"]]).strip().lower()
                hometown = str(row[header_index["BIRTH PLACE"]] or "").strip()

                with transaction.atomic():
                    emp = Employee(
                        username=username,
                        first_name=first,
                        middle_name=middle,
                        last_name=last,
                        visa_type=visa_type,
                        department=department,
                        designation=designation,
                        hire_date=hire_date,
                        primary_contact=tp_number,
                        hometown=hometown,
                        gender=gender_raw,
                        nationality=country_code,
                        special_notes=special_notes,
                        is_active=True,
                        is_staff=False,
                        is_superuser=False,
                    )
                    emp.set_password("carhub123")
                    emp.save()
                    emp.groups.set([group])
                    results["created"] += 1
            except Exception as exc:
                results["errors"].append(f"{sheet_name} row {row_num}: {exc}")
    return results


def import_company_visa_workbook(file: IO[bytes]) -> Dict[str, object]:
    """Import company visa employees from an Excel workbook.

    The workbook contains one sheet per branch. Each sheet includes employees of
    that branch with columns:

    NAME, DOJ, [DESIGNATION], TRADE LIC, WPS NO, NATIONALITY
    """

    wb = load_workbook(file, data_only=True)
    group, _ = Group.objects.get_or_create(name=EMPLOYEE_ROLE)
    results: Dict[str, object] = {"created": 0, "errors": []}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        headers = [str(cell.value).strip().upper() if cell.value else "" for cell in ws[1]]
        header_index = {h: i for i, h in enumerate(headers)}
        branch = Branch.objects.filter(name__iexact=sheet_name).first()
        if not branch:
            results["errors"].append(f"{sheet_name}: unknown branch")
            continue
        department, _ = Department.objects.get_or_create(branch=branch, name="General")

        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            try:
                name = str(row[header_index["NAME"]]).strip()
                first, middle, last = _split_name(name)

                designation = None
                if "DESIGNATION" in header_index:
                    designation_name = str(row[header_index["DESIGNATION"]] or "").strip()
                    if designation_name:
                        designation, _ = Designation.objects.get_or_create(
                            company=branch.company, name=designation_name
                        )

                trade_lic_no = str(row[header_index["TRADE LIC"]]).strip()
                if not trade_lic_no:
                    raise ValueError("missing trade license")
                trade_license, _ = TradeLicense.objects.get_or_create(
                    license_no=trade_lic_no,
                    defaults={"company": branch.company, "max_visas": 1},
                )
                trade_license.branches.add(branch)

                nationality_raw = str(row[header_index["NATIONALITY"]] or "").strip().upper()
                country_code = NATIONALITY_MAP.get(nationality_raw)
                if not country_code:
                    raise ValueError(f"Unsupported nationality '{nationality_raw}'")

                hire_raw = row[header_index["DOJ"]]
                hire_date: date | None = None
                if hire_raw:
                    if isinstance(hire_raw, datetime):
                        hire_date = hire_raw.date()
                    elif isinstance(hire_raw, date):
                        hire_date = hire_raw
                    else:
                        try:
                            hire_date = datetime.strptime(str(hire_raw), "%d %b %Y").date()
                        except ValueError:
                            hire_date = datetime.strptime(str(hire_raw), "%Y/%m/%d").date()

                wps_id = str(row[header_index["WPS NO"]]).strip()
                username_base = "".join(name.split()).lower()
                username = username_base
                suffix = 1
                while Employee.objects.filter(username=username).exists():
                    username = f"{username_base}{suffix}"
                    suffix += 1

                with transaction.atomic():
                    emp = Employee(
                        username=username,
                        first_name=first,
                        middle_name=middle,
                        last_name=last,
                        visa_type="company",
                        payment_status="wps",
                        wps_id=wps_id,
                        trade_license=trade_license,
                        department=department,
                        designation=designation,
                        hire_date=hire_date,
                        nationality=country_code,
                        employment_type="permanent",
                        is_active=True,
                        is_staff=False,
                        is_superuser=False,
                    )
                    emp.set_password("carhub123")
                    emp.save()
                    emp.groups.set([group])
                    results["created"] += 1
            except Exception as exc:
                results["errors"].append(f"{sheet_name} row {row_num}: {exc}")
    return results


def import_main_format_workbook(file: IO[bytes]) -> Dict[str, object]:
    """Import employees using the standardized main Excel format."""

    wb = load_workbook(file, data_only=True)
    ws = wb.active
    results: Dict[str, object] = {"created": 0, "errors": []}

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        results["errors"].append("Workbook is missing a header row")
        return results

    header_index: Dict[str, int] = {}
    for idx, raw_header in enumerate(header_row):
        canonical = MAIN_FORMAT_HEADER_MAP.get(_normalize_header(raw_header))
        if canonical and canonical not in header_index:
            header_index[canonical] = idx

    missing_columns = [col for col in MAIN_FORMAT_REQUIRED_COLUMNS if col not in header_index]
    if missing_columns:
        results["errors"].append(
            "Missing required columns: " + ", ".join(sorted(missing_columns))
        )
        return results

    group, _ = Group.objects.get_or_create(name=EMPLOYEE_ROLE)

    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(_is_empty(cell) for cell in row):
            continue
        try:
            first_raw = _stringify(row[header_index["first_name"]]).strip()
            if not first_raw:
                raise ValueError("missing first name")
            first = first_raw.title()

            middle = ""
            if "middle_name" in header_index:
                middle = _stringify(row[header_index["middle_name"]]).strip().title()

            last = ""
            if "last_name" in header_index:
                last = _stringify(row[header_index["last_name"]]).strip().title()

            branch_name = _stringify(row[header_index["branch"]]).strip()
            if not branch_name:
                raise ValueError("missing branch name")
            branch = Branch.objects.filter(name__iexact=branch_name).select_related("company").first()
            if not branch:
                raise ValueError(f"Unknown branch '{branch_name}'")
            department, _ = Department.objects.get_or_create(branch=branch, name="General")

            visa_raw = _stringify(row[header_index["visa_type"]]).strip()
            if not visa_raw:
                raise ValueError("missing visa type")
            visa_type = VISA_TYPE_MAP.get(visa_raw.replace(" ", "").upper())
            if not visa_type:
                raise ValueError(f"Unsupported visa type '{visa_raw}'")

            payment_raw = _stringify(row[header_index["payment_status"]]).strip()
            if not payment_raw:
                raise ValueError("missing payment status")
            payment_status = PAYMENT_STATUS_MAP.get(payment_raw.replace(" ", "").upper())
            if not payment_status:
                raise ValueError(f"Unsupported payment status '{payment_raw}'")
            if visa_type != "company" and payment_status != "cash":
                raise ValueError("Personal or visit visa requires cash payment")

            trade_license_number = ""
            if "trade_license" in header_index:
                trade_license_number = _stringify(row[header_index["trade_license"]]).strip()

            trade_license = None
            if visa_type == "company":
                if not trade_license_number:
                    raise ValueError("missing trade license number for company visa")
                trade_license = TradeLicense.objects.filter(
                    license_no__iexact=trade_license_number
                ).select_related("company").first()
                if not trade_license:
                    raise ValueError(f"Unknown trade license '{trade_license_number}'")
                if trade_license.company_id != branch.company_id:
                    raise ValueError(
                        f"Trade license '{trade_license_number}' belongs to a different company"
                    )
                trade_license.branches.add(branch)
            else:
                trade_license = None

            designation = None
            if "designation" in header_index:
                designation_name = _stringify(row[header_index["designation"]]).strip()
                if designation_name:
                    designation = Designation.objects.filter(
                        company=branch.company, name__iexact=designation_name
                    ).first()
                    if not designation:
                        designation = Designation.objects.create(
                            company=branch.company, name=designation_name
                        )

            wps_account_number = ""
            if "wps_account_number" in header_index:
                wps_account_number = _stringify(
                    row[header_index["wps_account_number"]]
                ).strip()
            if payment_status == "wps" and not wps_account_number:
                raise ValueError("WPS account number is required when payment status is WPS")
            if payment_status != "wps":
                wps_account_number = ""

            c3_id = ""
            if "c3_number" in header_index:
                c3_id = _stringify(row[header_index["c3_number"]]).strip()
            if visa_type != "company":
                c3_id = ""

            hire_date = _parse_hire_date(row[header_index["hire_date"]])

            gender = ""
            if "gender" in header_index:
                gender = _normalize_gender(row[header_index["gender"]])

            raw_username_parts = [part for part in [first, middle, last] if part]
            cleaned_username_parts = []
            for part in raw_username_parts:
                part_no_whitespace = "".join(part.split())
                cleaned_part = re.sub(r"[^A-Za-z0-9@.+_-]", "", part_no_whitespace)
                if cleaned_part:
                    cleaned_username_parts.append(cleaned_part)
            username_base = "".join(part.lower() for part in cleaned_username_parts)
            if not username_base:
                username_base = f"employee{row_num}"
            username = username_base
            suffix = 1
            while Employee.objects.filter(username=username).exists():
                username = f"{username_base}{suffix}"
                suffix += 1

            with transaction.atomic():
                employee = Employee(
                    username=username,
                    first_name=first,
                    middle_name=middle,
                    last_name=last,
                    visa_type=visa_type,
                    payment_status=payment_status,
                    wps_account_number=wps_account_number,
                    department=department,
                    designation=designation,
                    trade_license=trade_license,
                    hire_date=hire_date,
                    employment_type="permanent",
                    c3_id=c3_id,
                    gender=gender,
                    is_active=True,
                    is_staff=False,
                    is_superuser=False,
                )
                employee.set_password("carhub123")
                employee.full_clean()
                employee.save()
                employee.groups.set([group])
                results["created"] += 1
        except ValidationError as exc:
            if hasattr(exc, "message_dict"):
                parts = []
                for field, messages in exc.message_dict.items():
                    parts.append(f"{field}: {', '.join(messages)}")
                message = "; ".join(parts)
            else:
                message = "; ".join(exc.messages)
            results["errors"].append(f"Row {row_num}: {message}")
        except Exception as exc:
            results["errors"].append(f"Row {row_num}: {exc}")

    return results
